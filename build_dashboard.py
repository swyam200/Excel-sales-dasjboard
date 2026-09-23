import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter

df = pd.read_csv("retail_sales.csv", parse_dates=["Date"])
df["Month"] = df["Date"].dt.strftime("%Y-%m")

wb = Workbook()

# ---------- Sheet 1: Raw Data ----------
ws_data = wb.active
ws_data.title = "Raw Data"
headers = ["OrderID","Date","Region","Category","UnitsSold","UnitPrice","Revenue","Discount","NetRevenue","Month"]
ws_data.append(headers)
for c in range(1, len(headers)+1):
    cell = ws_data.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF", name="Arial")
    cell.fill = PatternFill("solid", fgColor="1F3864")
    cell.alignment = Alignment(horizontal="center")

for _, row in df.iterrows():
    ws_data.append([
        row["OrderID"], row["Date"].strftime("%Y-%m-%d"), row["Region"], row["Category"],
        row["UnitsSold"], row["UnitPrice"], row["Revenue"], row["Discount"], row["NetRevenue"], row["Month"]
    ])

for i, w in enumerate([10,12,10,12,10,10,12,10,12,10], start=1):
    ws_data.column_dimensions[get_column_letter(i)].width = w

last_row = ws_data.max_row

# ---------- Sheet 2: Region Summary (formulas, not hardcoded) ----------
ws_region = wb.create_sheet("Region Summary")
ws_region.append(["Region", "Total Revenue", "Orders", "Avg Order Value"])
for c in range(1,5):
    cell = ws_region.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF", name="Arial")
    cell.fill = PatternFill("solid", fgColor="1F3864")

regions = sorted(df["Region"].unique())
for i, r in enumerate(regions, start=2):
    ws_region.cell(row=i, column=1, value=r).font = Font(name="Arial")
    ws_region.cell(row=i, column=2, value=f"=ROUND(SUMIFS('Raw Data'!I2:I{last_row},'Raw Data'!C2:C{last_row},A{i}),2)").font = Font(name="Arial")
    ws_region.cell(row=i, column=3, value=f"=COUNTIFS('Raw Data'!C2:C{last_row},A{i})").font = Font(name="Arial")
    ws_region.cell(row=i, column=4, value=f"=ROUND(B{i}/C{i},2)").font = Font(name="Arial")
for i in range(1,5):
    ws_region.column_dimensions[get_column_letter(i)].width = 16

# Bar chart: revenue by region
bar = BarChart()
bar.title = "Revenue by Region"
bar.y_axis.title = "Revenue ($)"
data_ref = Reference(ws_region, min_col=2, min_row=1, max_row=1+len(regions))
cats_ref = Reference(ws_region, min_col=1, min_row=2, max_row=1+len(regions))
bar.add_data(data_ref, titles_from_data=True)
bar.set_categories(cats_ref)
bar.width, bar.height = 14, 8
ws_region.add_chart(bar, "F2")

# ---------- Sheet 3: Category Summary ----------
ws_cat = wb.create_sheet("Category Summary")
ws_cat.append(["Category", "Revenue", "% of Total"])
for c in range(1,4):
    cell = ws_cat.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF", name="Arial")
    cell.fill = PatternFill("solid", fgColor="1F3864")

cats = sorted(df["Category"].unique())
for i, cname in enumerate(cats, start=2):
    ws_cat.cell(row=i, column=1, value=cname).font = Font(name="Arial")
    ws_cat.cell(row=i, column=2, value=f"=ROUND(SUMIFS('Raw Data'!I2:I{last_row},'Raw Data'!D2:D{last_row},A{i}),2)").font = Font(name="Arial")
total_row = 2 + len(cats)
ws_cat.cell(row=total_row, column=1, value="Total").font = Font(bold=True, name="Arial")
ws_cat.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row-1})").font = Font(bold=True, name="Arial")
for i in range(2, total_row):
    ws_cat.cell(row=i, column=3, value=f"=ROUND(100*B{i}/$B${total_row},1)").font = Font(name="Arial")
for i in range(1,4):
    ws_cat.column_dimensions[get_column_letter(i)].width = 16

pie = PieChart()
pie.title = "Revenue Share by Category"
data_ref = Reference(ws_cat, min_col=2, min_row=1, max_row=total_row-1)
cats_ref = Reference(ws_cat, min_col=1, min_row=2, max_row=total_row-1)
pie.add_data(data_ref, titles_from_data=True)
pie.set_categories(cats_ref)
pie.width, pie.height = 12, 8
ws_cat.add_chart(pie, "E2")

# ---------- Sheet 4: Monthly Trend ----------
ws_month = wb.create_sheet("Monthly Trend")
ws_month.append(["Month", "Net Revenue"])
for c in range(1,3):
    cell = ws_month.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF", name="Arial")
    cell.fill = PatternFill("solid", fgColor="1F3864")

months = sorted(df["Month"].unique())
for i, m in enumerate(months, start=2):
    ws_month.cell(row=i, column=1, value=m).font = Font(name="Arial")
    ws_month.cell(row=i, column=2, value=f"=ROUND(SUMIFS('Raw Data'!I2:I{last_row},'Raw Data'!J2:J{last_row},A{i}),2)").font = Font(name="Arial")
for i in range(1,3):
    ws_month.column_dimensions[get_column_letter(i)].width = 14

line = LineChart()
line.title = "Monthly Net Revenue Trend"
line.y_axis.title = "Revenue ($)"
data_ref = Reference(ws_month, min_col=2, min_row=1, max_row=1+len(months))
cats_ref = Reference(ws_month, min_col=1, min_row=2, max_row=1+len(months))
line.add_data(data_ref, titles_from_data=True)
line.set_categories(cats_ref)
line.width, line.height = 18, 8
ws_month.add_chart(line, "D2")

wb.save("Sales_Dashboard.xlsx")
print("saved")
